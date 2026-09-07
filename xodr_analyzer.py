from __future__ import annotations

import argparse
import csv
import json
import os
import sys
import traceback
import zipfile
from datetime import datetime
from pathlib import Path
import xml.etree.ElementTree as ET

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def configure_console() -> None:
    if os.name == "nt":
        try:
            sys.stdout.reconfigure(encoding="utf-8")
            sys.stderr.reconfigure(encoding="utf-8")
        except Exception:
            pass


def to_float(value, default=0.0):
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def to_int(value, default=0):
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def safe_dataframe(records, columns):
    return pd.DataFrame(records, columns=columns)


def clean_for_excel(df: pd.DataFrame) -> pd.DataFrame:
    """
    Excel에 안전하게 저장하기 위해 복잡한 값과 제어문자를 정리한다.
    """
    cleaned = df.copy()

    def clean_value(value):
        if value is None:
            return None

        if isinstance(value, (dict, list, tuple, set)):
            value = json.dumps(value, ensure_ascii=False)

        if isinstance(value, str):
            # Excel XML에서 허용되지 않는 제어문자 제거
            value = "".join(
                ch for ch in value
                if ch in "\t\n\r" or ord(ch) >= 32
            )
            # Excel 셀 최대 길이
            if len(value) > 32767:
                value = value[:32767]

        return value

    for column in cleaned.columns:
        cleaned[column] = cleaned[column].map(clean_value)

    return cleaned


def autosize_and_style_excel(excel_path: Path) -> None:
    workbook = load_workbook(excel_path)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        max_rows_to_scan = min(worksheet.max_row, 1000)

        for column_index in range(1, worksheet.max_column + 1):
            max_length = 0

            for row_index in range(1, max_rows_to_scan + 1):
                value = worksheet.cell(row_index, column_index).value
                if value is None:
                    continue
                max_length = max(max_length, len(str(value)))

            adjusted_width = min(max(max_length + 2, 10), 45)
            worksheet.column_dimensions[
                get_column_letter(column_index)
            ].width = adjusted_width

    workbook.save(excel_path)


def validate_xlsx(excel_path: Path) -> None:
    """
    XLSX가 실제 ZIP 구조를 갖고 있고 openpyxl로 다시 열리는지 확인한다.
    """
    if not excel_path.exists():
        raise RuntimeError("Excel 파일이 생성되지 않았습니다.")

    if excel_path.stat().st_size < 1000:
        raise RuntimeError("Excel 파일 크기가 비정상적으로 작습니다.")

    if not zipfile.is_zipfile(excel_path):
        raise RuntimeError("생성된 파일이 정상적인 XLSX ZIP 구조가 아닙니다.")

    with zipfile.ZipFile(excel_path, "r") as archive:
        required_items = {
            "[Content_Types].xml",
            "xl/workbook.xml",
        }
        archive_names = set(archive.namelist())
        missing = required_items - archive_names
        if missing:
            raise RuntimeError(
                f"Excel 내부 필수 파일 누락: {sorted(missing)}"
            )

    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    workbook.close()


def build_dataframes(root: ET.Element):
    # Roads
    road_rows = []

    for road in root.findall("road"):
        road_link = road.find("link")
        predecessor_id = None
        predecessor_type = None
        successor_id = None
        successor_type = None

        if road_link is not None:
            predecessor = road_link.find("predecessor")
            successor = road_link.find("successor")

            if predecessor is not None:
                predecessor_id = predecessor.get("elementId")
                predecessor_type = predecessor.get("elementType")

            if successor is not None:
                successor_id = successor.get("elementId")
                successor_type = successor.get("elementType")

        road_rows.append({
            "road_id": road.get("id"),
            "road_name": road.get("name"),
            "length_m": to_float(road.get("length")),
            "junction_id": road.get("junction"),
            "rule": road.get("rule"),
            "predecessor_id": predecessor_id,
            "predecessor_type": predecessor_type,
            "successor_id": successor_id,
            "successor_type": successor_type,
        })

    roads_df = pd.DataFrame(road_rows)

    # Lanes
    lane_rows = []

    for road in root.findall("road"):
        road_id = road.get("id")

        for section_index, lane_section in enumerate(
            road.findall("./lanes/laneSection")
        ):
            section_s = to_float(lane_section.get("s"))

            for side_name in ("left", "center", "right"):
                side = lane_section.find(side_name)
                if side is None:
                    continue

                for lane in side.findall("lane"):
                    lane_link = lane.find("link")
                    predecessor_lane_id = None
                    successor_lane_id = None

                    if lane_link is not None:
                        predecessor = lane_link.find("predecessor")
                        successor = lane_link.find("successor")

                        if predecessor is not None:
                            predecessor_lane_id = predecessor.get("id")
                        if successor is not None:
                            successor_lane_id = successor.get("id")

                    width_elements = lane.findall("width")
                    first_width = width_elements[0] if width_elements else None

                    lane_rows.append({
                        "road_id": road_id,
                        "lane_section_index": section_index,
                        "section_s_m": section_s,
                        "side": side_name,
                        "lane_id": to_int(lane.get("id")),
                        "lane_type": lane.get("type"),
                        "level": lane.get("level"),
                        "predecessor_lane_id": predecessor_lane_id,
                        "successor_lane_id": successor_lane_id,
                        "width_record_count": len(width_elements),
                        "width_a": (
                            to_float(first_width.get("a"))
                            if first_width is not None
                            else None
                        ),
                    })

    lane_columns = [
        "road_id", "lane_section_index", "section_s_m", "side",
        "lane_id", "lane_type", "level", "predecessor_lane_id",
        "successor_lane_id", "width_record_count", "width_a"
    ]
    lanes_df = safe_dataframe(lane_rows, lane_columns)

    # Geometry
    geometry_rows = []
    geometry_tags = ("line", "arc", "spiral", "poly3", "paramPoly3")

    for road in root.findall("road"):
        road_id = road.get("id")

        for geometry_index, geometry in enumerate(
            road.findall("./planView/geometry")
        ):
            geometry_type = "unknown"
            curvature = None
            curv_start = None
            curv_end = None

            for tag_name in geometry_tags:
                shape = geometry.find(tag_name)
                if shape is None:
                    continue

                geometry_type = tag_name

                if tag_name == "arc":
                    curvature = to_float(shape.get("curvature"))
                elif tag_name == "spiral":
                    curv_start = to_float(shape.get("curvStart"))
                    curv_end = to_float(shape.get("curvEnd"))
                break

            geometry_rows.append({
                "road_id": road_id,
                "geometry_index": geometry_index,
                "s_m": to_float(geometry.get("s")),
                "x_m": to_float(geometry.get("x")),
                "y_m": to_float(geometry.get("y")),
                "hdg_rad": to_float(geometry.get("hdg")),
                "length_m": to_float(geometry.get("length")),
                "geometry_type": geometry_type,
                "curvature": curvature,
                "curv_start": curv_start,
                "curv_end": curv_end,
            })

    geometry_columns = [
        "road_id", "geometry_index", "s_m", "x_m", "y_m",
        "hdg_rad", "length_m", "geometry_type", "curvature",
        "curv_start", "curv_end"
    ]
    geometry_df = safe_dataframe(geometry_rows, geometry_columns)

    # Junctions
    junction_rows = []

    for junction in root.findall("junction"):
        base_junction = {
            "junction_id": junction.get("id"),
            "junction_name": junction.get("name"),
            "junction_type": junction.get("type"),
        }

        for connection in junction.findall("connection"):
            base_connection = {
                **base_junction,
                "connection_id": connection.get("id"),
                "incoming_road": connection.get("incomingRoad"),
                "connecting_road": connection.get("connectingRoad"),
                "linked_road": connection.get("linkedRoad"),
                "contact_point": connection.get("contactPoint"),
            }

            lane_links = connection.findall("laneLink")

            if not lane_links:
                junction_rows.append({
                    **base_connection,
                    "from_lane": None,
                    "to_lane": None,
                })

            for lane_link in lane_links:
                junction_rows.append({
                    **base_connection,
                    "from_lane": lane_link.get("from"),
                    "to_lane": lane_link.get("to"),
                })

    junction_columns = [
        "junction_id", "junction_name", "junction_type", "connection_id",
        "incoming_road", "connecting_road", "linked_road",
        "contact_point", "from_lane", "to_lane"
    ]
    junction_df = safe_dataframe(junction_rows, junction_columns)

    # Signals
    signal_rows = []

    for road in root.findall("road"):
        road_id = road.get("id")

        for signal in road.findall("./signals/signal"):
            signal_rows.append({
                "road_id": road_id,
                "signal_id": signal.get("id"),
                "signal_name": signal.get("name"),
                "s_m": to_float(signal.get("s")),
                "t_m": to_float(signal.get("t")),
                "z_offset_m": to_float(signal.get("zOffset")),
                "dynamic": signal.get("dynamic"),
                "orientation": signal.get("orientation"),
                "country": signal.get("country"),
                "type": signal.get("type"),
                "subtype": signal.get("subtype"),
                "value": signal.get("value"),
                "unit": signal.get("unit"),
                "height_m": to_float(signal.get("height")),
                "width_m": to_float(signal.get("width")),
            })

    signal_columns = [
        "road_id", "signal_id", "signal_name", "s_m", "t_m",
        "z_offset_m", "dynamic", "orientation", "country", "type",
        "subtype", "value", "unit", "height_m", "width_m"
    ]
    signals_df = safe_dataframe(signal_rows, signal_columns)

    # Objects
    object_rows = []

    for road in root.findall("road"):
        road_id = road.get("id")

        for obj in road.findall("./objects/object"):
            object_rows.append({
                "road_id": road_id,
                "object_id": obj.get("id"),
                "object_name": obj.get("name"),
                "object_type": obj.get("type"),
                "s_m": to_float(obj.get("s")),
                "t_m": to_float(obj.get("t")),
                "z_offset_m": to_float(obj.get("zOffset")),
                "hdg_rad": to_float(obj.get("hdg")),
                "pitch_rad": to_float(obj.get("pitch")),
                "roll_rad": to_float(obj.get("roll")),
                "length_m": to_float(obj.get("length")),
                "width_m": to_float(obj.get("width")),
                "height_m": to_float(obj.get("height")),
                "radius_m": to_float(obj.get("radius")),
                "orientation": obj.get("orientation"),
            })

    object_columns = [
        "road_id", "object_id", "object_name", "object_type", "s_m",
        "t_m", "z_offset_m", "hdg_rad", "pitch_rad", "roll_rad",
        "length_m", "width_m", "height_m", "radius_m", "orientation"
    ]
    objects_df = safe_dataframe(object_rows, object_columns)

    # Summary
    summary_df = roads_df.copy()

    if not lanes_df.empty:
        summary_df["lane_count"] = summary_df["road_id"].map(
            lanes_df.groupby("road_id").size()
        )
        summary_df["driving_lane_count"] = summary_df["road_id"].map(
            lanes_df[lanes_df["lane_type"] == "driving"]
            .groupby("road_id")
            .size()
        )
        summary_df["lane_section_count"] = summary_df["road_id"].map(
            lanes_df.groupby("road_id")["lane_section_index"].nunique()
        )
    else:
        summary_df["lane_count"] = 0
        summary_df["driving_lane_count"] = 0
        summary_df["lane_section_count"] = 0

    if not geometry_df.empty:
        summary_df["geometry_count"] = summary_df["road_id"].map(
            geometry_df.groupby("road_id").size()
        )
        summary_df["geometry_types"] = summary_df["road_id"].map(
            geometry_df.groupby("road_id")["geometry_type"].apply(
                lambda values: ", ".join(pd.unique(values))
            )
        )
    else:
        summary_df["geometry_count"] = 0
        summary_df["geometry_types"] = "없음"

    if not signals_df.empty:
        summary_df["signal_count"] = summary_df["road_id"].map(
            signals_df.groupby("road_id").size()
        )
    else:
        summary_df["signal_count"] = 0

    if not objects_df.empty:
        summary_df["object_count"] = summary_df["road_id"].map(
            objects_df.groupby("road_id").size()
        )
    else:
        summary_df["object_count"] = 0

    count_columns = [
        "lane_count", "driving_lane_count", "lane_section_count",
        "geometry_count", "signal_count", "object_count"
    ]
    summary_df[count_columns] = (
        summary_df[count_columns].fillna(0).astype(int)
    )
    summary_df["geometry_types"] = summary_df[
        "geometry_types"
    ].fillna("없음")

    return {
        "Roads": roads_df,
        "Lanes": lanes_df,
        "Geometry": geometry_df,
        "Junctions": junction_df,
        "Signals": signals_df,
        "Objects": objects_df,
        "Road Summary": summary_df,
    }


def build_quality_checks(dataframes):
    roads_df = dataframes["Roads"]
    lanes_df = dataframes["Lanes"]
    geometry_df = dataframes["Geometry"]
    junction_df = dataframes["Junctions"]
    summary_df = dataframes["Road Summary"]

    short_roads_df = (
        summary_df[summary_df["length_m"] < 10]
        .sort_values("length_m")
        .copy()
    )

    no_driving_lane_df = summary_df[
        (summary_df["junction_id"].astype(str) == "-1")
        & (summary_df["driving_lane_count"] == 0)
    ].copy()

    geometry_length_sum = (
        geometry_df.groupby("road_id")["length_m"].sum()
        if not geometry_df.empty
        else pd.Series(dtype=float)
    )

    geometry_length_check_df = roads_df[
        ["road_id", "length_m"]
    ].copy()

    geometry_length_check_df["geometry_length_sum_m"] = (
        geometry_length_check_df["road_id"]
        .map(geometry_length_sum)
        .fillna(0)
    )

    geometry_length_check_df["difference_m"] = (
        geometry_length_check_df["length_m"]
        - geometry_length_check_df["geometry_length_sum_m"]
    ).abs()

    geometry_length_errors_df = (
        geometry_length_check_df[
            geometry_length_check_df["difference_m"] > 0.01
        ]
        .sort_values("difference_m", ascending=False)
        .copy()
    )

    if not junction_df.empty:
        missing_lane_link_df = junction_df[
            junction_df["from_lane"].isna()
            | junction_df["to_lane"].isna()
        ].copy()
    else:
        missing_lane_link_df = junction_df.copy()

    duplicate_road_id_df = roads_df[
        roads_df.duplicated("road_id", keep=False)
    ].copy()

    duplicate_lane_id_df = lanes_df[
        lanes_df.duplicated(
            [
                "road_id",
                "lane_section_index",
                "side",
                "lane_id",
            ],
            keep=False,
        )
    ].copy()

    return {
        "Short Roads": short_roads_df,
        "No Driving Lane": no_driving_lane_df,
        "Geometry Length Check": geometry_length_check_df,
        "Geometry Length Errors": geometry_length_errors_df,
        "Missing LaneLink": missing_lane_link_df,
        "Duplicate Road IDs": duplicate_road_id_df,
        "Duplicate Lane IDs": duplicate_lane_id_df,
    }


def build_statistics(xodr_path, dataframes, quality_checks):
    roads_df = dataframes["Roads"]
    lanes_df = dataframes["Lanes"]
    geometry_df = dataframes["Geometry"]
    junction_df = dataframes["Junctions"]
    signals_df = dataframes["Signals"]
    objects_df = dataframes["Objects"]

    normal_roads = roads_df["junction_id"].astype(str).eq("-1").sum()
    junction_roads = roads_df["junction_id"].astype(str).ne("-1").sum()

    stats = {
        "파일명": xodr_path.name,
        "분석 시각": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "전체 road 수": len(roads_df),
        "전체 road 길이(m)": roads_df["length_m"].sum(),
        "평균 road 길이(m)": roads_df["length_m"].mean(),
        "최단 road 길이(m)": roads_df["length_m"].min(),
        "최장 road 길이(m)": roads_df["length_m"].max(),
        "일반 road 수": int(normal_roads),
        "교차로 내부 road 수": int(junction_roads),
        "전체 lane 행 수": len(lanes_df),
        "driving lane 행 수": (
            int(lanes_df["lane_type"].eq("driving").sum())
            if not lanes_df.empty
            else 0
        ),
        "geometry 수": len(geometry_df),
        "junction 수": (
            junction_df["junction_id"].nunique()
            if not junction_df.empty
            else 0
        ),
        "signal 수": len(signals_df),
        "object 수": len(objects_df),
        "10m 미만 road 수": len(quality_checks["Short Roads"]),
        "주행 차선 없는 일반 road 수": len(
            quality_checks["No Driving Lane"]
        ),
        "geometry 길이 불일치 수": len(
            quality_checks["Geometry Length Errors"]
        ),
        "laneLink 확인 대상 수": len(
            quality_checks["Missing LaneLink"]
        ),
        "중복 road ID 수": len(
            quality_checks["Duplicate Road IDs"]
        ),
        "중복 lane ID 수": len(
            quality_checks["Duplicate Lane IDs"]
        ),
    }

    return pd.DataFrame(stats.items(), columns=["항목", "값"])


def save_csv_backups(output_folder, all_dataframes):
    csv_folder = output_folder / "csv"
    csv_folder.mkdir(parents=True, exist_ok=True)

    for name, df in all_dataframes.items():
        safe_name = (
            name.replace(" ", "_")
            .replace("/", "_")
            .replace("\\", "_")
        )
        clean_for_excel(df).to_csv(
            csv_folder / f"{safe_name}.csv",
            index=False,
            encoding="utf-8-sig",
        )


def save_charts(output_folder, dataframes):
    charts_folder = output_folder / "charts"
    charts_folder.mkdir(parents=True, exist_ok=True)

    roads_df = dataframes["Roads"]
    lanes_df = dataframes["Lanes"]
    geometry_df = dataframes["Geometry"]
    junction_df = dataframes["Junctions"]
    summary_df = dataframes["Road Summary"]

    if not roads_df.empty:
        plt.figure(figsize=(10, 6))
        roads_df["length_m"].plot.hist(bins=40)
        plt.title("Road Length Distribution")
        plt.xlabel("Road length (m)")
        plt.ylabel("Road count")
        plt.tight_layout()
        plt.savefig(
            charts_folder / "01_road_length_distribution.png",
            dpi=200,
        )
        plt.close()

    if not lanes_df.empty:
        plt.figure(figsize=(10, 6))
        lanes_df["lane_type"].fillna("unknown").value_counts().plot.bar()
        plt.title("Lane Type Count")
        plt.xlabel("Lane type")
        plt.ylabel("Count")
        plt.xticks(rotation=45, ha="right")
        plt.tight_layout()
        plt.savefig(charts_folder / "02_lane_type_count.png", dpi=200)
        plt.close()

    if not geometry_df.empty:
        plt.figure(figsize=(10, 6))
        geometry_df[
            "geometry_type"
        ].fillna("unknown").value_counts().plot.bar()
        plt.title("Geometry Type Count")
        plt.xlabel("Geometry type")
        plt.ylabel("Count")
        plt.xticks(rotation=45, ha="right")
        plt.tight_layout()
        plt.savefig(
            charts_folder / "03_geometry_type_count.png",
            dpi=200,
        )
        plt.close()

    if not summary_df.empty:
        plt.figure(figsize=(10, 6))
        summary_df[
            "driving_lane_count"
        ].value_counts().sort_index().plot.bar()
        plt.title("Driving Lane Count per Road")
        plt.xlabel("Driving lane count")
        plt.ylabel("Road count")
        plt.tight_layout()
        plt.savefig(
            charts_folder / "04_driving_lane_distribution.png",
            dpi=200,
        )
        plt.close()

        top_roads = roads_df.nlargest(
            15, "length_m"
        ).sort_values("length_m")

        plt.figure(figsize=(10, 7))
        plt.barh(
            top_roads["road_id"].astype(str),
            top_roads["length_m"],
        )
        plt.title("Top 15 Longest Roads")
        plt.xlabel("Road length (m)")
        plt.ylabel("Road ID")
        plt.tight_layout()
        plt.savefig(
            charts_folder / "05_top15_longest_roads.png",
            dpi=200,
        )
        plt.close()

    if not junction_df.empty:
        junction_complexity = (
            junction_df.groupby("junction_id")
            .agg(
                connection_count=("connection_id", "nunique"),
                lane_link_count=("from_lane", "count"),
            )
            .sort_values("lane_link_count", ascending=False)
            .head(15)
            .sort_values("lane_link_count")
        )

        plt.figure(figsize=(10, 7))
        plt.barh(
            junction_complexity.index.astype(str),
            junction_complexity["lane_link_count"],
        )
        plt.title("Top 15 Complex Junctions")
        plt.xlabel("LaneLink count")
        plt.ylabel("Junction ID")
        plt.tight_layout()
        plt.savefig(
            charts_folder / "06_top15_complex_junctions.png",
            dpi=200,
        )
        plt.close()


def save_excel_atomic(excel_path, all_dataframes):
    """
    임시 XLSX에 먼저 저장하고 검증이 끝난 뒤 최종 파일명으로 교체한다.
    중간 오류가 나도 손상된 최종 XLSX가 남지 않는다.
    """
    temp_path = excel_path.with_name(
        f"{excel_path.stem}.tmp.xlsx"
    )

    if temp_path.exists():
        temp_path.unlink()

    if excel_path.exists():
        try:
            excel_path.unlink()
        except PermissionError as error:
            raise PermissionError(
                "기존 Excel 파일이 열려 있습니다. "
                "Excel을 닫고 다시 실행하세요."
            ) from error

    try:
        with pd.ExcelWriter(
            temp_path,
            engine="openpyxl",
            mode="w",
        ) as writer:
            for sheet_name, df in all_dataframes.items():
                safe_sheet_name = sheet_name[:31]
                clean_for_excel(df).to_excel(
                    writer,
                    sheet_name=safe_sheet_name,
                    index=False,
                )

        autosize_and_style_excel(temp_path)
        validate_xlsx(temp_path)
        temp_path.replace(excel_path)
        validate_xlsx(excel_path)

    except Exception:
        if temp_path.exists():
            temp_path.unlink()
        raise


def analyze_xodr(xodr_path: Path, output_folder: Path):
    output_folder.mkdir(parents=True, exist_ok=True)
    log_path = output_folder / "analysis.log"
    error_log_path = output_folder / "error.log"

    def log(message):
        timestamp = datetime.now().strftime("%H:%M:%S")
        line = f"[{timestamp}] {message}"
        print(line)
        with log_path.open("a", encoding="utf-8") as file:
            file.write(line + "\n")

    if log_path.exists():
        log_path.unlink()
    if error_log_path.exists():
        error_log_path.unlink()

    try:
        log(f"입력 파일: {xodr_path}")

        if not xodr_path.exists():
            raise FileNotFoundError(
                f"XODR 파일을 찾을 수 없습니다: {xodr_path}"
            )

        if xodr_path.suffix.lower() != ".xodr":
            log("주의: 파일 확장자가 .xodr가 아닙니다.")

        log("XML 읽는 중")
        tree = ET.parse(xodr_path)
        root = tree.getroot()

        if root.tag != "OpenDRIVE":
            raise ValueError(
                f"최상위 태그가 OpenDRIVE가 아닙니다: {root.tag}"
            )

        log("DataFrame 생성 중")
        dataframes = build_dataframes(root)

        log("품질 검사 중")
        quality_checks = build_quality_checks(dataframes)

        log("기본 통계 생성 중")
        statistics_df = build_statistics(
            xodr_path,
            dataframes,
            quality_checks,
        )

        all_dataframes = {
            "Basic Statistics": statistics_df,
            **dataframes,
            **quality_checks,
        }

        log("CSV 백업 저장 중")
        save_csv_backups(output_folder, all_dataframes)

        log("그래프 저장 중")
        save_charts(output_folder, dataframes)

        excel_path = output_folder / (
            f"{xodr_path.stem}_XODR_Analysis.xlsx"
        )

        log("Excel 저장 및 무결성 검사 중")
        save_excel_atomic(excel_path, all_dataframes)

        log("분석 완료")
        log(f"Excel: {excel_path}")
        return excel_path

    except Exception as error:
        with error_log_path.open("w", encoding="utf-8") as file:
            file.write("XODR Analyzer Error Report\n")
            file.write("=" * 70 + "\n")
            file.write(f"시간: {datetime.now()}\n")
            file.write(f"입력 파일: {xodr_path}\n")
            file.write(f"오류: {error}\n\n")
            file.write(traceback.format_exc())

        print(f"\n[오류] {error}")
        print(f"오류 로그: {error_log_path}")
        raise


def parse_args():
    parser = argparse.ArgumentParser(
        description="XODR 자동 분석기 v2"
    )
    parser.add_argument(
        "xodr",
        help="분석할 .xodr 파일 경로",
    )
    parser.add_argument(
        "-o",
        "--output",
        help="결과 저장 폴더",
    )
    return parser.parse_args()


def main():
    configure_console()
    args = parse_args()

    xodr_path = Path(args.xodr).expanduser().resolve()

    if args.output:
        output_folder = Path(args.output).expanduser().resolve()
    else:
        output_folder = (
            xodr_path.parent
            / f"{xodr_path.stem}_analysis_result"
        )

    try:
        analyze_xodr(xodr_path, output_folder)
    except ET.ParseError as error:
        print(f"[XML 오류] {error}")
        sys.exit(2)
    except Exception:
        sys.exit(1)


if __name__ == "__main__":
    main()
